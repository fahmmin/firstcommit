"""Excel/Tally ledger import — owner drops their existing .xlsx, rows become invoices.

MVP-thin but REAL: openpyxl reads the sheet, we map a handful of known column
headers to invoice fields and write real store rows. Unmappable rows are counted
in `skipped`, never crash the request (AYUSH.md §5.1).
"""
from __future__ import annotations

import uuid
from datetime import date, datetime, timezone

from .. import deps

# header synonyms → canonical field. Matched case-insensitively on a normalized
# (lowercased, non-alnum-stripped) header cell.
_COLUMN_MAP = {
    "invoiceno": "invoice_no", "invoice": "invoice_no", "invoicenumber": "invoice_no",
    "billno": "invoice_no", "bill": "invoice_no",
    "buyer": "buyer", "customer": "buyer", "party": "buyer", "client": "buyer",
    "partyname": "buyer", "customername": "buyer",
    "amount": "amount", "total": "amount", "value": "amount", "totalamount": "amount",
    "grandtotal": "amount",
    "duedate": "due_date", "due": "due_date", "date": "due_date", "paymentdue": "due_date",
    "gst": "gst", "tax": "gst", "gstamount": "gst", "taxamount": "gst",
    "items": "items", "item": "items", "description": "items", "particulars": "items",
}


def _norm(s) -> str:
    return "".join(ch for ch in str(s or "").lower() if ch.isalnum())


def _to_amount(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    cleaned = str(v).replace("₹", "").replace(",", "").replace("Rs", "").replace("rs", "").strip()
    try:
        return float(cleaned)
    except ValueError:
        return None


def _to_date(v) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%d-%b-%Y", "%d %b %Y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return s or None  # keep raw string rather than dropping the row


def _days_overdue(due: str | None) -> int:
    if not due:
        return 0
    try:
        return max(0, (date.today() - date.fromisoformat(due)).days)
    except ValueError:
        return 0


def import_excel_impl(tenant_id: str, file_path: str, filename: str = "ledger.xlsx") -> dict:
    """Parse an .xlsx ledger → invoice rows. Returns contract-shaped summary."""
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)

    # first non-empty row is the header
    header = None
    for r in rows:
        if r and any(c is not None and str(c).strip() for c in r):
            header = r
            break
    col_field: dict[int, str] = {}
    if header:
        for idx, cell in enumerate(header):
            field = _COLUMN_MAP.get(_norm(cell))
            if field and field not in col_field.values():
                col_field[idx] = field

    imported, skipped, sample = 0, 0, []
    have_amount_col = "amount" in col_field.values()
    for r in rows:
        if not r or not any(c is not None and str(c).strip() for c in r):
            continue
        rec: dict = {}
        for idx, field in col_field.items():
            if idx < len(r):
                rec[field] = r[idx]
        amount = _to_amount(rec.get("amount")) if have_amount_col else None
        buyer = str(rec.get("buyer") or "").strip()
        if amount is None or amount <= 0 or not buyer:
            skipped += 1
            continue
        due_date = _to_date(rec.get("due_date"))
        gst = _to_amount(rec.get("gst"))
        inv = {
            "id": f"inv-{uuid.uuid4().hex[:6]}",
            "invoice_no": str(rec.get("invoice_no") or f"INV-{uuid.uuid4().hex[:4].upper()}").strip(),
            "buyer": buyer,
            "amount": amount,
            "items": str(rec.get("items") or "").strip(),
            "issue_date": date.today().isoformat(),
            "due_date": due_date or "",
            "status": "overdue" if _days_overdue(due_date) > 0 else "sent",
            "days_overdue": _days_overdue(due_date),
            "source": "excel_import",
        }
        if gst is not None:
            inv["gst"] = gst
        deps.store.put_invoice(tenant_id, inv)
        imported += 1
        if len(sample) < 3:
            sample.append({"invoice_no": inv["invoice_no"], "buyer": inv["buyer"],
                           "amount": inv["amount"], "due_date": inv["due_date"]})
    wb.close()

    file_id = f"f-imp-{uuid.uuid4().hex[:6]}"
    if imported:
        deps.record_action("excel_imported", {"imported": imported, "skipped": skipped,
                                              "collection": "invoices"})
        deps.log_activity(tenant_id, "excel_imported",
                          f"Imported {imported} invoices from {filename}"
                          + (f" ({skipped} rows skipped)" if skipped else ""))
    return {
        "file_id": file_id, "filename": filename, "collection": "invoices",
        "imported": imported, "skipped": skipped, "sample": sample,
    }
